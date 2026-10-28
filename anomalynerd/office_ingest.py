"""Word (.docx) and Excel (.xlsx) input for AnomalyNerd.

Both formats store tables as real cell grids, so extraction is exact: no layout guessing, no
merged-with-prose problem. Each table (a Word table, or an Excel worksheet) is passed through the
same quality gate as PDF tables (pdf_ingest._assess_and_tidy, with structured=True), so tables
that are too small, non-numeric, or mis-structured are reported as not analyzed rather than
silently dropped or analyzed anyway. structured=True only skips the two PDF-specific checks for
figures misread as tables, which cannot occur in a real cell grid.

Word body text is also scanned for reported statistics (t, F, chi-square, p-values), exactly like
PDF body text. Excel has no prose, so it contributes tables only.

Requires python-docx (for .docx) and openpyxl (for .xlsx). Each is imported only when its format
is used, and a missing library is reported as an honest note instead of a crash.
"""
from .pdf_ingest import PdfParse, _assess_and_tidy, _extract_stat_strings


def _cell_text(value) -> str:
    """Render a cell value as the text a reader would see. Floats that are whole numbers are
    printed without a trailing '.0' so integers stay integers; None becomes blank."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _trim_grid(rows):
    """Drop fully blank rows and trailing fully blank columns (spreadsheets often carry empty
    formatted cells around the real data)."""
    rows = [r for r in rows if any(c != "" for c in r)]
    if not rows:
        return rows
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    while width > 0 and all(r[width - 1] == "" for r in rows):
        width -= 1
        rows = [r[:width] for r in rows]
    # drop leading fully blank columns too
    while rows and rows[0] and all(r[0] == "" for r in rows):
        rows = [r[1:] for r in rows]
    return rows


def _finalize(path, raw_tables, stat_strings, notes, n_pages=0, text_pages=0):
    """Run each raw grid through the shared quality gate and build a PdfParse."""
    tables = []
    for idx, (page, grid) in enumerate(raw_tables):
        et = _assess_and_tidy(grid, page, idx, structured=True)
        et.source_path = path
        tables.append(et)
    if not tables:
        notes.append("no tables were found in this file")
    uniq = list(dict.fromkeys(stat_strings))
    return PdfParse(path=path, n_pages=n_pages, tables=tables, stat_strings=uniq,
                    text_pages=text_pages, notes=notes)


def parse_docx(path: str) -> PdfParse:
    """Parse a Word document: every table becomes a candidate results table, and paragraph text
    is scanned for reported statistics."""
    notes = []
    try:
        import docx
    except ImportError:
        return PdfParse(path=path, n_pages=0, tables=[], stat_strings=[],
                        notes=["python-docx is not installed, so Word tables could not be read; "
                               "`pip install python-docx` to enable .docx input"])
    doc = docx.Document(path)
    raw_tables = []
    for ti, table in enumerate(doc.tables, start=1):
        grid = []
        for row in table.rows:
            # python-docx repeats a merged cell once per spanned column; keep that, it mirrors
            # how the table reads visually and keeps every row the same width.
            grid.append([_cell_text(c.text) for c in row.cells])
        grid = _trim_grid(grid)
        if grid:
            raw_tables.append((ti, grid))   # "page" = table number, for readable locations
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    stats = _extract_stat_strings(text) if text else []
    return _finalize(path, raw_tables, stats, notes, n_pages=len(doc.tables),
                     text_pages=1 if text else 0)


def parse_xlsx(path: str) -> PdfParse:
    """Parse an Excel workbook: every non-empty worksheet becomes a candidate results table.
    Formulas are read as their last computed values (what the user sees)."""
    notes = []
    try:
        import openpyxl
    except ImportError:
        return PdfParse(path=path, n_pages=0, tables=[], stat_strings=[],
                        notes=["openpyxl is not installed, so Excel sheets could not be read; "
                               "`pip install openpyxl` to enable .xlsx input"])
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    raw_tables = []
    try:
        for si, ws in enumerate(wb.worksheets, start=1):
            grid = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            grid = _trim_grid(grid)
            if grid:
                raw_tables.append((si, grid))   # "page" = sheet number
            else:
                notes.append(f"sheet '{ws.title}' is empty")
    finally:
        wb.close()
    return _finalize(path, raw_tables, [], notes, n_pages=len(raw_tables))
